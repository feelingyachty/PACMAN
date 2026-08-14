#!/usr/bin/env python3
"""Build the Feeling Yachty Koray/Tongberg SEO audit Word document."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path("/workspace/docs/seo/audits/Feeling-Yachty-SEO-Framework-Audit-2026-08-14.docx")
NAVY = RGBColor(0x14, 0x16, 0x0F)
GOLD = RGBColor(0xC4, 0x92, 0x00)
MUTED = RGBColor(0x5C, 0x57, 0x40)
ALERT = RGBColor(0xB4, 0x3A, 0x1F)


def set_run(run, *, size=11, bold=False, color=NAVY, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def add_p(doc, text, *, size=11, bold=False, color=NAVY, italic=False, space_after=8, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color, italic=italic)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = NAVY if level > 1 else GOLD
        run.font.name = "Calibri"
    return p


def bullet(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    if bold_lead:
        r = p.add_run(bold_lead)
        set_run(r, bold=True)
        r = p.add_run(text)
        set_run(r)
    else:
        r = p.add_run(text)
        set_run(r)
    return p


def shade_header(cell, hex_color="14160F"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {qn("w:fill"): hex_color, qn("w:val"): "clear"})
    tcPr.append(shd)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        set_run(run, size=10, bold=True, color=RGBColor(0xFF, 0xF6, 0xD8))
        shade_header(cell)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            set_run(run, size=10)
    doc.add_paragraph()
    return table


def score_row(name, score, note):
    return [name, f"{score} / 5", note]


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.95)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run(
        "Feeling Yachty  ·  Confidential SEO audit  ·  PACMAN / Megaman  ·  14 August 2026  ·  Page "
    )
    set_run(fr, size=8, color=MUTED)
    # page number field
    from docx.oxml import OxmlElement

    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    fr2 = footer.add_run()
    set_run(fr2, size=8, color=MUTED)
    fr2._r.append(fld1)
    fr2._r.append(instr)
    fr2._r.append(fld2)

    add_p(doc, "PACMAN  ·  MEGAMAN", size=12, bold=True, color=GOLD, space_after=2)
    add_p(
        doc,
        "Feeling Yachty SEO Framework Audit",
        size=26,
        bold=True,
        space_after=4,
    )
    add_p(
        doc,
        "Koray Tuğberk GÜBÜR / Corey Tongberg Holistic SEO  ·  feelingyachty.com",
        size=13,
        italic=True,
        color=MUTED,
        space_after=12,
    )
    add_p(
        doc,
        "Prepared for Fernando  ·  14 August 2026  ·  Live site + sitemap_index.xml",
        size=11,
        color=MUTED,
        space_after=6,
    )
    add_p(
        doc,
        "This is an operator audit, not a traffic forecast. No rankings, sessions, or revenue "
        "figures are invented. Every claim below is tied to a live URL, a sitemap count, or a "
        "search snapshot taken the same day.",
        size=11,
        italic=True,
        space_after=16,
    )

    heading(doc, "1. Verdict", 1)
    add_p(
        doc,
        "The site is a strong commercial catalog that is a weak semantic content network. "
        "Feeling Yachty already has the hard parts of charter SEO: real inventory, visible prices, "
        "WhatsApp and GHL booking paths, reviews, and two real destinations. The Koray / Tongberg "
        "framework would fail the site on the next test: one source context, one owner URL per "
        "query cluster, consistent facts, and a map that does not drown money pages in near-duplicates.",
    )
    add_p(
        doc,
        "In plain language: Google (and a guest) has to work too hard to decide which URL is "
        "“the” Miami yacht charter page, which phone is real, whether fuel is included, and "
        "whether this brand is Miami-only or Miami + Panama. That is cost of retrieval. That is "
        "the opposite of topical authority.",
    )

    add_table(
        doc,
        ["Lens", "Score", "What I saw"],
        [
            score_row("Source context", 2, "Homepage H1 is Miami Yacht Rental. Panama is a footnote. Brand is not the roof."),
            score_row("EAV completeness (listings)", 4, "Cards and Panama routes expose size, capacity, hours, price. Strong."),
            score_row("EAV completeness (money hubs)", 2, "Hubs repeat the same inventory widget instead of owning one job."),
            score_row("Intent / format match", 3, "Catalog + filters match DO intent. KNOW answers are late or duplicated."),
            score_row("Information gain", 3, "Birthday itinerary tip is excellent. Most hubs do not add unique facts."),
            score_row("Internal bridges", 2, "Many folders, same boats, weak ownership. Tips often loop to the same hub."),
            score_row("Cannibalization safety", 1, "Home, /miami-yacht-rental/, /miami-yacht-charters/, cheap, affordable, luxury…"),
            score_row("Trust / identity / NAP", 1, "Phones, review counts, inclusions, and city labels contradict each other."),
        ],
    )
    add_p(
        doc,
        "Overall: 2.3 / 5 against the PACMAN Koray bar. Fix contradictions and ownership before "
        "publishing more tips. More nodes will make this worse.",
        bold=True,
    )

    heading(doc, "2. Method and limits", 1)
    bullet(doc, "Sitemap source: https://feelingyachty.com/sitemap_index.xml (28 child sitemaps).")
    bullet(doc, "Parsed 1,721 unique <loc> URLs on 14 August 2026 (861 English, 860 Spanish /es/).")
    bullet(
        doc,
        "Scraped live pages including home, Miami rental, Miami charters, Panama rentals, cheap, "
        "affordable, Fort Lauderdale, contact, crew, reviews, pink rentals, book, a vessel listing "
        "(/fleet/miami/26ft-bayliner-fendi/), a tip (/miami-yacht-tips/birthday-yacht-itinerary-miami/), "
        "/test/, and /es/.",
    )
    bullet(
        doc,
        "Brand and head-term search snapshots via Firecrawl (not Google Search Console). A snapshot "
        "is not a ranking history.",
    )
    bullet(
        doc,
        "Framework: skills/seo-corey-tongberg (public Holistic SEO / Topical Authority materials by "
        "Koray Tuğberk GÜBÜR, operationalized for Feeling Yachty). Playbooks A–E and C scorecards.",
    )
    bullet(
        doc,
        "Cloudflare blocked raw datacenter curl. Pages were read through a browser scrape. Utility "
        "pages (cart, checkout) were confirmed in the sitemap, not fully shop-tested.",
    )
    add_p(
        doc,
        "I did not have GSC, GBP admin, or a full crawl of all 1,721 bodies. I did not invent "
        "which pages “rank.” I did inventory the graph and read the URLs that decide source context.",
        italic=True,
    )

    heading(doc, "3. Source context — the first failure", 1)
    add_p(
        doc,
        "Koray’s first rule: lock source context. Feeling Yachty’s operating truth (from your own "
        "stack and the founder line on the homepage) is: a private yacht charter / experience brand "
        "serving Miami and Panama, monetizing bookings and qualified leads.",
    )
    add_p(
        doc,
        "The live homepage does not say that. The H1 is “Miami Yacht Rental.” The first screen is "
        "Biscayne Bay, birthdays, pink yachts, and a Miami inventory widget (“Showing 27 of 178 yachts”). "
        "Panama appears much later, in a founder bio: Fernando Yemail, licensed Florida yacht broker, "
        "“serving both Miami and Panama.” A crawler that only reads the top of / will classify this "
        "site as a Miami rental marketplace, not a two-city charter operator.",
    )
    add_p(
        doc,
        "Spanish /es/ is a Miami homepage translation (“Alquiler de yates en Miami”), not a "
        "Panama-aware brand roof. hreflang is implemented (en-US, es-CO, en, es) — that part is fine. "
        "The problem is you mirrored the wrong roof.",
    )
    heading(doc, "What “good” looks like", 2)
    bullet(doc, "Homepage owns brand + two destinations + how booking works. It does not compete with the Miami money URL.")
    bullet(doc, "/miami-yacht-rental/ (or one chosen slug) owns Miami commercial charter.")
    bullet(doc, "/panama-yacht-rentals/ owns Panama commercial charter — this page is already the closest to the bar.")
    bullet(doc, "Fort Lauderdale is either a real third money node or a Miami-area location page. It cannot be both.")

    heading(doc, "4. Sitemap inventory (facts)", 1)
    add_p(doc, "Child sitemaps in the index (Yoast-style WordPress + custom post types + WooCommerce):")
    add_table(
        doc,
        ["Child sitemap", "URLs in file", "What it is"],
        [
            ["product-sitemap.xml", "592", "Woo products — yachts and catering SKUs"],
            ["miami_yacht_tips-sitemap.xml", "248", "Miami tips / auto-blog layer"],
            ["panama_yacht_tips-sitemap.xml", "192", "Panama tips / auto-blog layer"],
            ["page-sitemap.xml", "152", "WordPress pages (hubs, utilities, tests)"],
            ["msy_rentals-sitemap.xml", "74", "Miami super-yacht rentals CPT"],
            ["miami-mega-yachts-sitemap.xml", "56", "Mega yacht CPT"],
            ["miami-yacht-location-sitemap.xml", "48", "Location / neighborhood pages"],
            ["miami-yacht-specials-sitemap.xml", "44", "Specials / cheap-folder listings"],
            ["ft-lauderdale-events-sitemap.xml", "42", "Fort Lauderdale events"],
            ["miami-yacht-venue-sitemap.xml", "40", "Event-venue yachts"],
            ["video-sitemap.xml", "38", "Video attachments"],
            ["category-sitemap.xml", "34", "Blog / fleet categories"],
            ["panama_yacht_cpt-sitemap.xml", "28", "Panama yacht CPT"],
            ["miami-yacht-party-sitemap.xml", "28", "Occasion / party CPT"],
            ["ftlauderdale-charter-sitemap.xml", "26", "FTL charter CPT"],
            ["miami-yacht-charters-sitemap.xml", "26", "Charter CPT"],
            ["miami-yacht-services-sitemap.xml", "22", "DJ, catering, photo, chef…"],
            ["miami-water-sports-sitemap.xml", "18", "Water sports"],
            ["product_cat-sitemap.xml", "14", "Woo categories"],
            ["Other CPTs / tips / local", "—", "Pink yacht, menus, tours, maintenance, KML"],
        ],
    )
    add_p(doc, "English URL folder counts (unique EN paths):")
    add_table(
        doc,
        ["Folder / stem", "EN URLs", "Framework note"],
        [
            ["/fleet/…", "296", "Real vessels + catering products. Prices baked into slugs."],
            ["/miami-yacht-tips/…", "125", "Outer section. Homepage cards labeled “automation”."],
            ["/panama-yacht-tips/…", "97", "Outer section. Same risk as Miami tips."],
            ["/miami-super-yacht-rentals/…", "37", "Overlaps mega / luxury / 75ft+ hubs."],
            ["/miami-mega-yachts/…", "28", "Same entity type, second folder."],
            ["/miami-yacht-location/…", "24", "Valid attribute expansion if facts stay local."],
            ["/cheapest-yacht-rentals/…", "22", "Third copy of many Miami boats."],
            ["/ft-lauderdale-events/…", "21", "City identity must stay Fort Lauderdale."],
            ["/miami-yacht-venue/…", "20", "Occasion/venue vs rental cannibalization risk."],
            ["/panama-yacht-rentals/…", "16", "Core. Keep and reinforce."],
            ["/miami-yacht-party/…", "15", "Occasion cluster — assign one owner."],
            ["/miami-yacht-charters/…", "14", "Near-duplicate of /miami-yacht-rental/."],
        ],
    )
    add_p(
        doc,
        "1,721 indexable URLs is not automatically “topical authority.” Koray’s cost-of-retrieval "
        "warning applies: a newsfirehose or SKU firehose raises crawl cost without raising "
        "competence. Catering SKUs such as "
        "/fleet/miami-catering/fettuccine-alfredo-350-serves-13-guests-italian-catering-by-feeling-yachty/ "
        "are products, not topical nodes. Panama vessels sitting under /fleet/miami-catering/ "
        "(Ocean Phoenix, Parker Monaco) are folder-level identity errors.",
    )

    heading(doc, "5. Cannibalization registry", 1)
    add_p(
        doc,
        "Playbook A requires one owner URL per query cluster. Today several clusters have three to "
        "six plausible owners. I am assigning owners as a recommendation, not as a claim about current rankings.",
    )
    add_table(
        doc,
        ["Query cluster", "Recommended owner", "Current competitors (live)", "Decision"],
        [
            [
                "Brand / Feeling Yachty",
                "https://feelingyachty.com/  (rewritten)",
                "Home currently targets Miami rental. Reviews, crew, contact also speak for the brand.",
                "Rewrite home. Do not keep Miami H1 on /.",
            ],
            [
                "Miami yacht rental / charter (head)",
                "/miami-yacht-rental/",
                "/ , /miami-yacht-charters/, /miami-yacht-rentals-directory/, emoji slug …feeling-yachty-2/, /luxury-yacht-rentals/, /miami-yachts-booking/",
                "Pick one. 301 or noindex the rest of the hubs.",
            ],
            [
                "Cheap / budget Miami charter",
                "/cheap-yacht-rentals-miami/  OR merge into owner with a filter",
                "/affordable-yachts-rentals-miami/ is a near-clone (same “all fees included” pitch + same widget).",
                "Keep one modifier page or none. Do not keep both.",
            ],
            [
                "Pink yacht Miami",
                "/miami-pink-yacht-rentals/",
                "Home filters, party CPT, individual pink listings, cheapest-folder copies.",
                "Hub + listings. Hub must not repeat the full 400-yacht grid.",
            ],
            [
                "Occasion: birthday Miami",
                "/miami-yacht-party/birthday-yacht-rentals/  or the tip if you make the tip the KNOW owner",
                "Party CPT + /miami-yacht-tips/birthday-yacht-itinerary-miami/ + 20/30/50-person pages.",
                "One DO hub, one KNOW guide. Cross-link. Different H1s.",
            ],
            [
                "Capacity modifiers (20 / 30 / 50 / 100 person)",
                "One capacity guide + filtered inventory",
                "/20-person-yacht-rental-miami/, /30-person-…, /50-person-…, /l00-person-yacht-rentals/ (typo).",
                "Fix typo. Do not grow a page per integer.",
            ],
            [
                "Fort Lauderdale charter",
                "/fort-lauderdale-yacht-rentals/",
                "Page H1 says Fort Lauderdale, next H1 says “Best North Miami Yacht Rentals.” Location CPT + events CPT.",
                "Rewrite or it will rank for the wrong city — or neither.",
            ],
            [
                "Panama yacht rental",
                "/panama-yacht-rentals/",
                "Directory, party, sales, maintenance, tours, 97 tips. Cleaner than Miami.",
                "Keep. Do not clone Miami’s hub sprawl.",
            ],
            [
                "Single vessel (example: 26ft Bayliner Fendi)",
                "/fleet/miami/26ft-bayliner-fendi/",
                "Also in cheapest-yacht-rentals/miami/26ft-pink-bayliner-… and inventory cards on 5+ hubs.",
                "One canonical listing. Others 301.",
            ],
        ],
    )

    heading(doc, "6. Page scorecards (Playbook C)", 1)
    heading(doc, "6.1 Homepage  ·  feelingyachty.com/", 2)
    add_p(
        doc,
        "Central entity claimed: Miami yacht rental — not Feeling Yachty the brand. "
        "Same pricing widget and FAQ pattern as the Miami hubs. Tips module is tagged "
        "“automation” in the HTML/text scrape (bachelor cost, bachelorette decorations). "
        "That is the n8n auto-blog leaking onto the money roof.",
    )
    add_table(
        doc,
        ["Dimension", "Score", "Note"],
        [
            score_row("Source context", 1, "Miami H1. Panama buried."),
            score_row("EAV", 3, "Widget has size/price/capacity. Brand EAV is thin."),
            score_row("Intent / format", 3, "Good DO catalog. Wrong URL for that job."),
            score_row("Information gain", 2, "Trust badges + founder line. No unique brand definition up top."),
            score_row("Bridges", 2, "Tips all point at /miami-yacht-tips/. Weak Panama bridge."),
            score_row("Cannibalization safety", 1, "Competes with /miami-yacht-rental/ and /miami-yacht-charters/."),
            score_row("Trust", 2, "2,500+ on the hero vs 2,400+ in the founder paragraph on the same page."),
        ],
    )
    add_p(doc, "Decision: rewrite. Keep /. Change the job.", bold=True)

    heading(doc, "6.2 Miami money pair  ·  /miami-yacht-rental/  and  /miami-yacht-charters/", 2)
    add_p(
        doc,
        "These are the same page with different lipstick. Both open with 400+ boats, the same "
        "WhatsApp script, the same GHL booking widget, the same “How Miami Yacht Rental Pricing Works” "
        "block, and the same filter UI. Counts on the scrape: rental hub “27 of 49 yachts”; charters "
        "hub “27 of 178 yachts”; homepage “27 of 178.” If those numbers are filters, the copy still "
        "says 400+ everywhere. That is an information-gap and a trust gap.",
    )
    add_p(doc, "Decision: one owner. 301 the loser. Do not “differentiate” with synonyms.", bold=True)

    heading(doc, "6.3 Cheap vs affordable", 2)
    add_p(
        doc,
        "/cheap-yacht-rentals-miami/ H1: “Cheap Yacht Rentals in Miami With All Fees Included.” "
        "/affordable-yachts-rentals-miami/ H1: “Affordable Yacht Rentals in Miami With All Fees Included.” "
        "Same four benefit blocks, same payment line (Klarna / Affirm / crypto), same inventory widget. "
        "Affordable adds “Every Charter Includes: Captain, Crew, Fuel…” and then the shared widget says "
        "“Crew and fuel are additional.” That is a contradiction on the same URL.",
    )
    add_p(doc, "Decision: merge. One budget page, or a filter on the owner hub. Never both.", bold=True)

    heading(doc, "6.4 Panama  ·  /panama-yacht-rentals/", 2)
    add_p(
        doc,
        "This is the best money page I read. Central entity is clear. Pricing is route-based "
        "(Panama Bay vs Taboga vs Contadora vs Colón) — that is real EAV, not synonym stuffing. "
        "Filters by vessel type, size, and group size. WhatsApp uses +507 202-1729. "
        "Keep this pattern. Do not let Panama tips and maintenance pages outgrow it.",
    )
    add_table(
        doc,
        ["Dimension", "Score", "Note"],
        [
            score_row("Source context", 4, "Panama charter operator. Distinct from Miami."),
            score_row("EAV", 4, "Route × duration × price on cards."),
            score_row("Intent / format", 4, "Catalog matches commercial-investigation."),
            score_row("Information gain", 3, "Route pricing is the unique gain. Deepen inclusions/exclusions."),
            score_row("Bridges", 3, "Needs tighter links from tips and tours back to this owner."),
            score_row("Cannibalization safety", 3, "Directory / party / sales exist but are not clones of this hub."),
            score_row("Trust", 3, "Phone disagrees with the contact page (see §8)."),
        ],
    )
    add_p(doc, "Decision: keep and reinforce. This is the template for Miami after the merge.", bold=True)

    heading(doc, "6.5 Fort Lauderdale  ·  /fort-lauderdale-yacht-rentals/", 2)
    add_p(
        doc,
        "First H1: Fort Lauderdale yacht rentals. Immediate second H1: “Best North Miami Yacht Rentals.” "
        "Body: Haulover, Sunny Isles, Bal Harbour, then links to Miami Beach and Fort Lauderdale. "
        "“We are the ONLY company in Miami that gives away A FREE HOUR.” Superlative + wrong city "
        "on a Fort Lauderdale slug. Typos (“Avaliable”). Older template than the Miami neon inventory.",
    )
    add_p(doc, "Decision: rewrite as a true FTL/Haulover location page, or 301 into Miami location architecture.", bold=True)

    heading(doc, "6.6 Vessel listing  ·  /fleet/miami/26ft-bayliner-fendi/", 2)
    add_p(
        doc,
        "Gallery-first listing with Woo booker. Attributes exist (duration, passenger count, deposit rule). "
        "Problems: the Miami rental hub card for the same boat says From $800 / 4 hours / max 13 guests. "
        "The listing shows $500.00 due today and a passenger selector that goes to 40. If 13 is the legal "
        "cap, a 40-person selector is a safety and trust defect, not an SEO nit. "
        "A cheapest-folder twin also exists. One vessel, one URL.",
    )
    add_p(doc, "Decision: keep /fleet/miami/{slug}/ as the listing pattern. 301 folder copies. Lock capacity to the real max.", bold=True)

    heading(doc, "6.7 Outer node that works  ·  /miami-yacht-tips/birthday-yacht-itinerary-miami/", 2)
    add_p(
        doc,
        "This is what Koray actually wants from outer content. Quick answer in the first screen. "
        "Duration, group size, boarding, itinerary templates (4-hour / 6-hour / sunset). Bridges to "
        "/30-person-yacht-rental-miami/ and /yacht-safety-protocols/. Operator tone, not encyclopedia tourism. "
        "Use this as the quality gate for every tip. Most of the 222 tip URLs will not pass it.",
    )
    add_p(doc, "Decision: keep. Make it the KNOW owner for birthday itinerary. Do not let a second birthday guide spawn.", bold=True)

    heading(doc, "6.8 Indexed junk", 2)
    add_p(doc, "These are in the sitemap. They should not be:")
    bullet(doc, "https://feelingyachty.com/test/ — H1 “TEST” plus a 227-yacht inventory clone.")
    bullet(doc, "https://feelingyachty.com/test-page/")
    bullet(doc, "https://feelingyachty.com/miami-sailing-tips/test/")
    bullet(doc, "https://feelingyachty.com/maintenance-tips/test-yacht-maintenance-tips/")
    bullet(doc, "https://feelingyachty.com/event-services-panama/test-event-services/")
    bullet(doc, "https://feelingyachty.com/panama-yacht-services/panama-yacht-services-test/")
    bullet(doc, "https://feelingyachty.com/cart/ and /checkout/ and /confirmed/")
    bullet(doc, "https://feelingyachty.com/elementor-52093/ and /12-2/ and /3-party-form/")
    bullet(doc, "https://feelingyachty.com/l00-person-yacht-rentals/ — typo of 100.")
    add_p(doc, "Decision: noindex + remove from sitemap this week. Tests are not “content.”", bold=True)

    heading(doc, "7. Brand identity and Brand SERP", 1)
    add_p(
        doc,
        "Playbook E: the brand query should resolve to a yacht charter operator with one NAP and "
        "one story. Snapshot for “Feeling Yachty” (Firecrawl search, 14 August 2026): official site, "
        "Instagram @feeling.yachty, Yelp, Tripadvisor, Trustpilot, Facebook. Panama is an emoji on "
        "Instagram, not a second Brand SERP estate. That is acceptable if Miami is the demand center. "
        "It is not acceptable if phones and addresses disagree.",
    )
    add_table(
        doc,
        ["Surface", "Identifier seen", "Conflict"],
        [
            ["Site CTAs (Miami)", "+1 (954) 246-3636", "Primary on-site Miami number"],
            ["Crew page SMS", "+1 (754) 325-3827", "Third US number on an identity page"],
            ["Yelp + Facebook snapshot", "(786) 352-8857", "Does not match the site hero number"],
            ["Yelp address snapshot", "3201 NW 24th Street Rd, Fl 2, Miami, FL 33142", "Must match GBP + footer + schema or be removed"],
            ["Panama rentals page", "+507 202-1729", "WhatsApp and tel on the money page"],
            ["Contact page Panama", "+507 202-1279", "Two digits swapped vs the money page"],
            ["Review counts on-site", "1,300+ / 2,400+ / 2,500+ / 2,700+", "Four different claims across FTL, home, cheap, crew"],
            ["Fleet size on-site", "400+ vs 178 vs 49 vs 31 vs 227", "Pick a definition (unique bookable vessels) and use it"],
        ],
    )
    add_p(
        doc,
        "Jason Barnard’s Brand SERP point (as Koray cites it): corroboration. Right now third-party "
        "profiles and on-site pages are teaching the entity different phone numbers. Fix the sources. "
        "Do not add more neighborhood GBPs until NAP is one row.",
    )
    add_p(
        doc,
        "Head-term snapshot: “private yacht charter Miami” returned directories and competitors "
        "(YachtCharterFleet, Tropicalboat, Miami Charters, Prime Luxury, Biscayne Lady, Getmyboat) "
        "in the first results of this pull. feelingyachty.com did not appear in that short snapshot. "
        "That is not a GSC report. It is a reminder that synonym hubs are not the same as winning the "
        "query network.",
        italic=True,
    )
    add_p(
        doc,
        "“yacht charter Panama” snapshot was international crewed-week directories and a Panama City Beach, "
        "Florida false friend. Your Panama page’s route-level EAV is the correct counter to that SERP. "
        "Use Taboga / Contadora / San Blas facts you can stand behind — do not copy Miami sandbar copy.",
    )

    heading(doc, "8. Contradiction hunt (must fix before new content)", 1)
    add_p(
        doc,
        "Koray: historical data compounds when facts stay stable. PACMAN upgrade: hunt contradictions "
        "before publish. These are live today.",
    )
    add_table(
        doc,
        ["Fact", "Version A", "Version B", "Risk"],
        [
            ["Inclusions", "Cheap/affordable/FTL: all fees / captain / fuel included", "Shared widget: “Crew and fuel are additional.”", "Guest + crawler distrust. Possible ads issue."],
            ["Deposit", "Only $1,400+ pays 20%", "Fendi listing: “Due today $500” plus 20% rule for $1,400+", "Needs one ruleset in writing."],
            ["Fendi price", "Hub card: $800 / 4 hours", "Listing: $500 due today", "Same entity, two prices."],
            ["Fendi capacity", "Hub: max 13", "Listing selector: 1–40", "Safety + legal. Fix even if SEO were perfect."],
            ["Reviews", "1,300+ / 2,400+ / 2,500+ / 2,700+", "Tripadvisor badge 2025 cited", "Pick one audited number or say “reviews across platforms” with no integer."],
            ["Panama phone", "202-1729", "202-1279", "Leads go to the wrong line."],
            ["US phones", "954 / 754 / 786", "Three public numbers", "Brand SERP fragmentation."],
            ["City", "FTL slug", "North Miami H1", "Wrong entity on the URL."],
            ["Superlative", "“ONLY company in Miami that gives away a free hour”", "Unverified on-page", "Remove or prove."],
        ],
    )

    heading(doc, "9. Topical map as it exists vs as it should", 1)
    heading(doc, "Core today (keep)", 2)
    bullet(doc, "Miami charter inventory (one hub + /fleet/miami/{vessel}/).")
    bullet(doc, "Panama charter inventory (keep /panama-yacht-rentals/ + Panama vessel CPTs).")
    bullet(doc, "Occasions that you actually sell: birthday, bachelorette, corporate, party — one hub each.")
    bullet(doc, "Yacht classes you actually sell: pink, sail, fishing, mega/superyacht — one hub each, not three.")
    bullet(doc, "Services that attach to a charter: catering, DJ, photo — as add-on hubs, not 100 SKU URLs if they dilute.")
    bullet(doc, "Identity: contact, crew/safety, reviews, charter agreement.")
    heading(doc, "Outer today (keep only with bridges)", 2)
    bullet(doc, "Tips that look like the birthday itinerary: process, routes, objections, what to bring.")
    bullet(doc, "Location pages with real marina/pickup facts (Miami Beach, River, Haulover) — not keyword mirrors.")
    bullet(doc, "Panama routes: Taboga, Contadora, Colón, San Blas — only if ops runs them.")
    heading(doc, "Reject / prune / noindex", 2)
    bullet(doc, "Test URLs, Elementor leftovers, cart/checkout, typo slugs.")
    bullet(doc, "Second and third copies of the same vessel in /cheapest-yacht-rentals/ and mega/super twins.")
    bullet(doc, "Tips that exist only because the auto-blog fired (homepage already labels them “automation”).")
    bullet(doc, "Panama maintenance / cleaning / visual-check pages unless you truly sell yacht care as a line of business with a roof context. If they are thin, they dilute charter identity.")
    bullet(doc, "Horseback riding Panama as a yacht-site node unless it is a real packaged add-on with a booking path.")

    heading(doc, "10. Auto-blogging and cost of retrieval", 1)
    add_p(
        doc,
        "n8n workflow “Feeling Yachty - WordPress Auto Blogging” (gKezuNipPn2PjqTV) is active in your "
        "ops docs. The homepage tip cards scrape with the word “automation.” Miami + Panama tip sitemaps "
        "were rewritten 14 August 2026 (same day as this audit). That is a publishing firehose.",
    )
    add_p(
        doc,
        "Koray’s SaaS / newsfirehose warning applies. 222 tip URLs can help if each one closes an "
        "information gap and bridges to a money owner. They hurt if they repeat “yacht rental Miami” "
        "with a new title. PACMAN rule already on file: n8n topics must be on the map. Until the map "
        "and ownership table exist, pause net-new tip publishing. Prune or noindex the tests. Keep "
        "the birthday itinerary class of post.",
    )

    heading(doc, "11. Spanish / hreflang", 1)
    add_p(
        doc,
        "Almost 1:1 EN/ES URL count (861 / 860) means you cloned the graph, including the problems. "
        "hreflang tags on the page sitemap are present. Good. Do not translate junk. When you 301 a "
        "duplicate EN hub, 301 its /es/ twin in the same change. Spanish home should follow the new "
        "brand roof, not “Alquiler de yates en Miami” as the corporate H1 unless Spanish demand is "
        "Miami-only — and even then the brand still needs a roof.",
    )

    heading(doc, "12. What is already working (do not break)", 1)
    bullet(doc, "Commercial machinery: WhatsApp deep links, click-to-call, GHL scheduling widget (LeadConnector).")
    bullet(doc, "Listing cards with duration × price tables. Panama route × price tables.")
    bullet(doc, "Filters (size, pink, free hour, budget) — these are attributes. Keep them on ONE hub.")
    bullet(doc, "Founder attribution and Tripadvisor / review surfaces — after the numbers are reconciled.")
    bullet(doc, "hreflang scaffolding.")
    bullet(doc, "The birthday itinerary tip as a model outer node.")
    bullet(doc, "Panama money page structure — copy this discipline to Miami after the merge.")

    heading(doc, "13. Priority queue (no calendar theater)", 1)
    add_p(doc, "Tier 0 — trust. Ship before any new blog node.", bold=True)
    bullet(doc, "One US phone, one Panama phone, one address set. Align site, GBP, Yelp, Facebook, schema.")
    bullet(doc, "One review sentence. Either an audited integer or no integer.")
    bullet(doc, "One inclusions rule. The widget and the H2 cannot disagree.")
    bullet(doc, "Noindex tests, cart, checkout, Elementor leftovers, /l00-person…")
    bullet(doc, "Lock listing capacity and price to ops truth (Fendi is the exhibit).")
    add_p(doc, "Tier 1 — ownership.", bold=True)
    bullet(doc, "Write the query → URL table. Approve it. Then 301.")
    bullet(doc, "Homepage becomes brand + Miami + Panama. Miami H1 moves off /.")
    bullet(doc, "Merge /miami-yacht-charters/ into /miami-yacht-rental/ (or the reverse — pick one slug and stop).")
    bullet(doc, "Merge cheap/affordable. Rewrite Fort Lauderdale or demote it.")
    bullet(doc, "One canonical per vessel. Kill cheapest-folder and CPT copies.")
    add_p(doc, "Tier 2 — map and outer.", bold=True)
    bullet(doc, "Freeze auto-blog until topics are on the map.")
    bullet(doc, "Keep/upgrade tips that pass the birthday-itinerary bar. Prune the rest.")
    bullet(doc, "Occasion hubs with different H1s than the Miami head term.")
    add_p(doc, "Tier 3 — reinforcement.", bold=True)
    bullet(doc, "Only after ownership is clean: deepen Miami EAV (inclusions matrix, departure marinas, occasion FAQs) on the single owner.")
    bullet(doc, "Mirror Panama’s route-style clarity onto Miami itineraries (sandbar vs skyline vs Key Biscayne) as attributes, not new money URLs.")

    heading(doc, "14. Open facts needed from ops (do not invent)", 1)
    bullet(doc, "Canonical US and Panama phone + WhatsApp + booking hours.")
    bullet(doc, "Legal passenger caps per vessel; who may exceed 13 and on which boats.")
    bullet(doc, "What “included” means: captain, crew, fuel, ice, dock fees, gratuity.")
    bullet(doc, "How many unique bookable vessels exist this week in Miami vs Panama — not “400+” until defined.")
    bullet(doc, "Which review integer you will stand behind, and which platforms it includes.")
    bullet(doc, "Whether Fort Lauderdale / Haulover is a real departure product or a Miami-area alias.")
    bullet(doc, "Whether Panama yacht maintenance / sales / horseback are real sold lines.")
    bullet(doc, "Whether every /es/ URL is a true translation or a stub.")

    heading(doc, "15. Megaman’s read", 1)
    add_p(
        doc,
        "Cori’s framework is the right judge for this site, and the site is currently using the "
        "wrong half of it. You have coverage. You do not have borders. You have entities (boats, "
        "cities, occasions). You do not have one document per entity type. You have historical "
        "publishing (tips updated today). You do not have historical consistency (phones, prices, "
        "review counts).",
    )
    add_p(
        doc,
        "If we apply the framework honestly, the next month of SEO work is subtraction and "
        "alignment, not another 40 posts. Corey should not get a “write more Miami tips” brief "
        "until the ownership table is approved. Pacman should verify the 301s and the NAP pass "
        "the way he verifies a meta change: live, not by word count.",
    )
    add_p(
        doc,
        "I can turn §13 into a Corey proposal (map + 301 list + homepage brief) when you want "
        "that gated. This document is the audit, not the implementation.",
        italic=True,
    )

    heading(doc, "Appendix A — Recommended owner URLs (draft)", 1)
    add_table(
        doc,
        ["Cluster", "Owner URL"],
        [
            ["Brand", "https://feelingyachty.com/"],
            ["Miami charter (head)", "https://feelingyachty.com/miami-yacht-rental/"],
            ["Panama charter (head)", "https://feelingyachty.com/panama-yacht-rentals/"],
            ["Contact / NAP", "https://feelingyachty.com/contact-feeling-yachty/"],
            ["Safety / crew", "https://feelingyachty.com/feeling-yachty-crew/"],
            ["Reviews", "https://feelingyachty.com/feeling-yachty-reviews/"],
            ["Pink Miami", "https://feelingyachty.com/miami-pink-yacht-rentals/"],
            ["Party / occasion Miami", "https://feelingyachty.com/miami-yacht-party/"],
            ["Birthday KNOW", "https://feelingyachty.com/miami-yacht-tips/birthday-yacht-itinerary-miami/"],
            ["Vessel pattern", "https://feelingyachty.com/fleet/miami/{vessel-slug}/"],
            ["Book / convert", "https://feelingyachty.com/book/  (or GHL only — do not let it rank as a money clone)"],
        ],
    )

    heading(doc, "Appendix B — Attribution", 1)
    add_p(
        doc,
        "Framework concepts: publicly published Holistic SEO / Topical Authority materials by "
        "Koray Tuğberk GÜBÜR (you call this Corey / Cori Tongberg). Brand SERP language also "
        "follows Jason Barnard as cited in those materials. This audit is PACMAN operationalization "
        "for Feeling Yachty. It does not use or claim Koray’s paywalled course. It does not forecast "
        "traffic.",
    )
    add_p(
        doc,
        "Primary evidence: feelingyachty.com sitemap_index.xml and the live URLs listed in §2, "
        "scraped 14 August 2026.",
    )

    doc.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
